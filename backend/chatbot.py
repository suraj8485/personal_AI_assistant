from pathlib import Path
from groq import Groq
from dotenv import load_dotenv
from time import sleep
from docx import Document
from pypdf import PdfReader
import os

from fastapi import FastAPI
from pydantic import BaseModel

app = FastAPI()

load_dotenv()

my_api_key = os.getenv("GROQ_API_KEY")

if not my_api_key:
    raise ValueError("where is your API key")

client = Groq(api_key = my_api_key)

model = "openai/gpt-oss-120b"

# say llm for extract my information from docs or pdf given.
def extract_details(reader):
    system_prompt = """
        you are a personal AI assistant and expert document Assistant. 
        # TASK :-
        extract the information from the document file A60205324095_suraj_sharma.docx
        take only those information which is mentioned in the document 
        do not add something at your own it is strictly prohibited.
        """
        
    user_prompt = f""" extract the information from the document which is {reader}"""

    return ASK_LLM(system_prompt,user_prompt)


# call the LLM for asking the question.
def ASK_LLM(system_prompt,user_prompt):

    system_message = {
        "role" : "system",
        "content" : system_prompt
    }
    message = {
        "role" : "user",
        "content" : user_prompt
    }
    messages = [system_message, message]
    response = client.chat.completions.create(model=model,messages = messages,temperature = 0.2)
    answer = response.choices[0].message.content
    
    return answer


# return the answer of the asked question about me
def ask_question(question : str, my_details):
    system_prompt = f""" 
        You are a friendly personal assistant. Answer questions using only the resume
        You are a friendly personal assistant. Answer questions using only the resume
        context provided in the conversation.
        
        Match questions by meaning, not just exact wording. For example, “Where are
        you from?” can be answered using a village, town, or place of origin stated
        in the resume.
        
        Answer in a clear, engaging, pointwise format:
        - Start with a short, relevant heading. You may use one suitable emoji.
        - Put each useful fact on its own line, starting with “•”.
        - bold the key points which question asked.
        - Keep the answer concise and easy to read.
        - For a simple question with one fact, give one concise bullet.
        - If the resume does not contain enough information, say:
          “I have no information about this.”
        - Do not invent or assume facts.
**Important:** `{my_details}` is the only source of truth.

"""

    #user_prompt = """what is the name of my sisters."""

    return ASK_LLM(system_prompt,question)


# read the pdf files.
def read_pdf(file_path):
    reader = PdfReader(file_path)
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()

        if page_text:
            text+= page_text + "\n"

    return text

# Read the word file.
def read_docx(file_path):
    document = Document(file_path)
    text = ""

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text+= paragraph.text + "\n"

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text+= cell.text + "\n"

    return text


def read_resume(file_path):
    if file_path.suffix.lower() == ".pdf":
        return read_pdf(file_path)

    if file_path.suffix.lower() == ".docx":
        return read_docx(file_path)
    else:
        return None


#class ChatRequest(BaseModel):
    question: str

# starter function fastapi automatically calls

my_details = None 
@app.on_event("startup")
def load_resume():

    global my_details

    my_details = read_docx(
        Path("A60205324095_suraj_sharma.docx")
    )
    
    print("Resume loaded successfully")


@app.get("/")
def home():
    return {
        "message" : "It is a Home Page"
    }


from fastapi.middleware.cors import CORSMiddleware

# Allow your frontend to talk to the backend
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],  # for development. Later change to your frontend URL
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

class ChatRequest(BaseModel):
    question : str

@app.post("/chat")
async def chat(request: ChatRequest):
    global my_details
    if my_details is None:
        return {
            "answer": "Resume is not loaded."
        }
    
    answer = ask_question(
        request.question,
         my_details
    )

    return {
        "answer": answer
    }


from fastapi.responses import FileResponse

@app.get("/download-resume")
async def download_resume():
    return FileResponse(
        path="A60205324095_suraj_sharma.docx",
        filename="A60205324095_suraj_sharma.docx",
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

#reader = read_resume(Path("A60205324095_suraj_sharma.docx"))
#print(extract_details(reader))
#my_details = extract_details(reader)
#sleep(5)
#result = ask_question(my_details)
#print(result)

